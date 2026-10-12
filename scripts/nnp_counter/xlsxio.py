"""อ่าน/เขียนไฟล์ xlsx (report, history, master, แท็บ missing) ด้วย openpyxl"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import HISTORY_COLUMNS, MISSING_HEADER, REPORT_COLUMNS, NnpError

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header_cell(c):
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center")
    c.border = BOX


def new_workbook(columns: list[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(columns)
    for c in ws[1][: len(columns)]:
        style_header_cell(c)
    return wb


def _num(v):
    """float ที่เป็นจำนวนเต็ม -> int เพื่อให้เทียบกันง่าย"""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def read_rows(path: Path, sheet: str | None = None, expect_header: list[str] | None = None) -> list[dict]:
    """อ่านตารางที่แถวแรกเป็นหัวคอลัมน์ -> list ของ dict (ข้ามแถวว่าง)"""
    path = Path(path)
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    if expect_header and header[: len(expect_header)] != expect_header:
        raise NnpError(f"หัวคอลัมน์ของ {path.name} ไม่ตรง: พบ {header[:len(expect_header)]} ต้องเป็น {expect_header}")
    out = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({header[i]: _num(r[i]) for i in range(min(len(header), len(r))) if header[i]})
    return out


def read_master(path: Path, cfg) -> list[dict]:
    """คืนรายการเครื่องตามลำดับใน master: {unit, machine, building}"""
    path = Path(path)
    if not path.exists():
        raise NnpError(f"ไม่พบไฟล์ master: {path}")
    rows = read_rows(path, cfg.master_sheet)
    cols = cfg.master_columns
    out = []
    for r in rows:
        machine = r.get(cols["machine"])
        if machine is None or str(machine).strip() == "":
            continue
        out.append({"unit": r.get(cols["unit"]), "machine": str(machine), "building": r.get(cols["building"])})
    return out


def _open(path: Path, columns: list[str]):
    path = Path(path)
    if path.exists():
        wb = load_workbook(path)
        ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
        header = [str(c.value).strip() if c.value is not None else "" for c in ws[1][: len(columns)]]
        if header != columns:
            raise NnpError(f"หัวคอลัมน์ของ {path.name} ไม่ตรง: พบ {header} ต้องเป็น {columns}")
    else:
        wb = new_workbook(columns)
        ws = wb["Sheet1"]
    return wb, ws


def _find_row(ws, col_idx: int, key: str):
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col_idx).value
        if v is not None and str(v).strip().upper() == key.upper():
            return r
    return None


def _first_empty_row(ws, col_idx: int) -> int:
    r = ws.max_row + 1
    while r > 2 and ws.cell(r - 1, col_idx).value in (None, ""):
        r -= 1
    return max(r, 2)


def upsert_report(path: Path, rows: list[dict], flagged: set[str], flag_color: str) -> dict:
    """เขียนแถวลง Sheet1 คีย์ = serialNumber; ระบายแดงแถวใน flagged"""
    wb, ws = _open(path, REPORT_COLUMNS)
    added = updated = 0
    fill = PatternFill("solid", start_color=flag_color, end_color=flag_color)
    for row in rows:
        r = _find_row(ws, 2, row["serialNumber"])
        if r is None:
            r = _first_empty_row(ws, 2)
            added += 1
        else:
            updated += 1
        for i, col in enumerate(REPORT_COLUMNS, start=1):
            v = row.get(col)
            c = ws.cell(r, i)
            c.value = "" if v is None else v
            if col == "remaining" and isinstance(v, (int, float)):
                c.number_format = "0%"
        if row["serialNumber"].upper() in flagged:
            for i in range(1, len(REPORT_COLUMNS) + 1):
                ws.cell(r, i).fill = fill
    wb.save(path)
    return {"added": added, "updated": updated}


def upsert_history(path: Path, rows: list[dict]) -> dict:
    wb, ws = _open(path, HISTORY_COLUMNS)
    added = updated = 0
    for row in rows:
        r = _find_row(ws, 1, row["Id"])
        if r is None:
            r = _first_empty_row(ws, 1)
            added += 1
        else:
            updated += 1
        for i, col in enumerate(HISTORY_COLUMNS, start=1):
            v = row.get(col)
            c = ws.cell(r, i)
            if col == "Timestamp" and isinstance(v, str):
                try:
                    v = dt.datetime.strptime(v, "%Y-%m-%d %H:%M:%S")
                except ValueError:
                    pass
            c.value = v
            if col == "Timestamp" and isinstance(v, dt.datetime):
                c.number_format = "yyyy-mm-dd h:mm:ss"
    wb.save(path)
    return {"added": added, "updated": updated}


def write_missing(path: Path, table: list[list], summary: list[list]) -> None:
    """เขียนแท็บ missing ใหม่ทั้งแท็บ (ล้างแล้วเขียน): ตาราง A1.. และสรุป F1:G4"""
    path = Path(path)
    wb = load_workbook(path)
    if "missing" in wb.sheetnames:
        ws = wb["missing"]
        ws.delete_rows(1, ws.max_row + 5)
        ws.delete_cols(1, ws.max_column + 5)
        for rng in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rng))
    else:
        ws = wb.create_sheet("missing")
    for r, row in enumerate(table, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c).value = v
    for r, row in enumerate(summary, start=1):
        for c, v in enumerate(row, start=6):
            ws.cell(r, c).value = v
    for c in range(1, len(MISSING_HEADER) + 1):
        style_header_cell(ws.cell(1, c))
    for r in range(1, len(summary) + 1):
        style_header_cell(ws.cell(r, 6))
    for col in range(1, 8):
        width = 8
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=True):
            v = row[0]
            if v is not None:
                width = max(width, min(len(str(v)) + 2, 80))
        ws.column_dimensions[get_column_letter(col)].width = width
    wb.save(path)


def read_missing_summary(path: Path) -> dict | None:
    path = Path(path)
    if not path.exists():
        return None
    wb = load_workbook(path, data_only=True)
    if "missing" not in wb.sheetnames:
        return None
    ws = wb["missing"]
    return {"total": _num(ws.cell(1, 7).value), "sent": _num(ws.cell(2, 7).value),
            "missing": _num(ws.cell(3, 7).value),
            "table_rows": sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value not in (None, ""))}
