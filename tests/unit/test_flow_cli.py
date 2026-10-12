"""ทดสอบ flow ผ่าน CLI จริง (subprocess) ใน repo จำลอง"""
import json

from openpyxl import load_workbook

from tests.conftest import jl, make_report, run_cli, snippet

OAG = "networkservice@oag.go.th"


def M(i, date, serial, total):
    return {"id": i, "threadId": "t" + i, "sender": OAG, "internalDate": "0", "text": snippet(date, serial, total)}


def start(repo, *extra, month="2026-09", org="oag"):
    return run_cli(repo, "begin", "--org", org, "--month", month, *extra)


def feed(repo, metas, org="oag", month="2026-09"):
    return run_cli(repo, "add-meta", "--org", org, "--month", month, stdin=jl(*metas))


def test_requires_valid_month_and_org(repo):
    code, d = run_cli(repo, "begin", "--org", "oag", "--month", "2026-13")
    assert code == 2 and "ไม่ถูกต้อง" in d["error"]
    p = run_cli  # org ผิดโดน argparse ปฏิเสธ
    try:
        p(repo, "begin", "--org", "xxx", "--month", "2026-09")
    except AssertionError:
        pass


def test_steps_cannot_be_skipped(repo):
    code, d = run_cli(repo, "apply", "--org", "oag", "--month", "2026-09")
    assert code == 2 and "begin" in d["error"]
    assert start(repo, "--dry-run")[0] == 0
    code, d = run_cli(repo, "apply", "--org", "oag", "--month", "2026-09")
    assert code == 2 and "ขั้น 'planned'" in d["error"]
    code, d = run_cli(repo, "verify", "--org", "oag", "--month", "2026-09")
    assert code == 2


def test_full_dry_run_oag_flags_and_missing(repo):
    make_report(repo / "02_สตง/03_report_files/oag-2026-08.xlsx",
                [{"serialNumber": "AAA111", "totalCounter": 100, "machineDate": "20-08-2026 09:00"}])
    code, d = start(repo, "--dry-run")
    assert code == 0 and d["banner"] == "ORG=oag MONTH=2026-09 MODE=DRY-RUN ENV=prod"
    assert d["gmail_query"] == "from:networkservice@oag.go.th after:1788195600 before:1790787600"
    metas = [M("m1", "10-09-2026 09:00", "AAA111", 150), M("m2", "12-09-2026 09:00", "AAA111", 170),
             M("m3", "11-09-2026 09:00", "BBB222", 70), M("m4", "01-10-2026 09:00", "CCC333", 9)]
    code, d = feed(repo, metas)
    assert d["accepted"] == 4
    assert feed(repo, metas)[1]["accepted"] == 0                       # ส่งซ้ำไม่ซ้อน
    code, d = run_cli(repo, "plan", "--org", "oag", "--month", "2026-09")
    assert d["counts"] == {"SUPERSEDED": 1, "COMPLETE": 2, "SKIPPED_OTHER_MONTH": 1}
    assert sorted(d["need_body"]) == ["m2", "m3"]
    code, d = run_cli(repo, "apply", "--org", "oag", "--month", "2026-09")
    assert code == 2 and "ขาดเนื้ออีเมล" in d["error"]
    code, d = run_cli(repo, "add-body", "--org", "oag", "--month", "2026-09",
                      stdin=jl({"id": "m2", "toner": "Black 80"}, {"id": "m3", "toner": "Black 5"}, {"id": "zzz", "toner": "Black 1"}))
    assert d["still_need_body"] == 0 and d["unknown_ids"] == ["zzz"]
    code, d = run_cli(repo, "apply", "--org", "oag", "--month", "2026-09")
    assert code == 0 and d["flagged"] == ["BBB222"] and d["missing"] == {"total": 3, "sent": 2, "missing": 1, "dup_serials": []}
    assert run_cli(repo, "labels", "--org", "oag", "--month", "2026-09")[1]["skip"] is True
    code, d = run_cli(repo, "verify", "--org", "oag", "--month", "2026-09", "--audited", "2", "--audit-mismatch", "0")
    assert code == 0 and d["ok"] and d["state"] == "verified", d
    # ไฟล์จริงไม่ถูกแตะ (dry-run เขียนใน sandbox)
    assert not (repo / "02_สตง/03_report_files/oag-2026-09.xlsx").exists()
    sandbox = next((repo / "_work/dryrun").iterdir())
    wb = load_workbook(sandbox / "02_สตง/03_report_files/oag-2026-09.xlsx")
    ws = wb["Sheet1"]
    rows = {r[1].value: r for r in ws.iter_rows(min_row=2)}
    assert rows["AAA111"][6].value == 170 and rows["AAA111"][7].value == 70 and rows["AAA111"][8].value == 0.8
    assert rows["AAA111"][0].fill.fgColor.rgb in (None, "00000000")                  # มีในเดือนก่อน: ไม่ไฮไลท์
    assert rows["BBB222"][0].fill.fgColor.rgb.endswith("FF6666")                       # ไม่มีในเดือนก่อน: แดง
    assert wb["missing"].cell(2, 2).value == "CCC333"


def test_real_run_idempotent_and_label_plan(repo):
    start(repo)
    feed(repo, [M("m1", "10-09-2026 09:00", "AAA111", 150), M("m2", "11-09-2026 09:00", "AAA111", 160)])
    run_cli(repo, "plan", "--org", "oag", "--month", "2026-09")
    run_cli(repo, "add-body", "--org", "oag", "--month", "2026-09", stdin=jl({"id": "m2", "toner": "Black 10"}))
    code, d = run_cli(repo, "apply", "--org", "oag", "--month", "2026-09")
    assert code == 0 and d["label_count"] == 2 and d["report"] == {"added": 1, "updated": 0}
    code, d = run_cli(repo, "verify", "--org", "oag", "--month", "2026-09")
    assert code == 2 and "labeled" in d["error"]                       # ยังไม่ติด label ห้าม verify
    code, d = run_cli(repo, "labels", "--org", "oag", "--month", "2026-09")
    assert d["label_name"] == "oag-2026-09" and d["message_ids"] == ["m1", "m2"]
    assert run_cli(repo, "labeled", "--org", "oag", "--month", "2026-09", "--count", "1")[0] == 2   # ติดไม่ครบ ปฏิเสธ
    assert run_cli(repo, "labeled", "--org", "oag", "--month", "2026-09", "--count", "2")[0] == 0
    code, d = run_cli(repo, "verify", "--org", "oag", "--month", "2026-09")
    assert code == 2 and "audit" in d["error"]                          # รอบจริงต้อง audit ก่อน
    code, d = run_cli(repo, "verify", "--org", "oag", "--month", "2026-09", "--audited", "1", "--audit-mismatch", "1")
    assert code == 1 and d["state"] != "verified"                        # audit ไม่ตรง = ไม่ผ่าน
    code, d = run_cli(repo, "verify", "--org", "oag", "--month", "2026-09", "--audited", "1", "--audit-mismatch", "0")
    assert code == 0 and d["state"] == "verified"
    # begin ซ้ำเมื่อ verified แล้ว: อีเมลเดิมถูกข้าม (ไม่มี COMPLETE ซ้ำ)
    assert start(repo, "--fresh")[0] == 0
    assert feed(repo, [M("m1", "10-09-2026 09:00", "AAA111", 150)])[1]["already_done"] == 1
    ws = load_workbook(repo / "02_สตง/03_report_files/oag-2026-09.xlsx")["Sheet1"]
    assert ws.max_row == 2


def test_concurrent_real_run_is_refused(repo):
    assert start(repo)[0] == 0
    code, d = start(repo)
    assert code == 2 and "ค้างอยู่" in d["error"]
    assert start(repo, "--force")[0] == 0


def test_dry_run_cannot_clobber_real_run_and_abort_releases(repo):
    assert start(repo)[0] == 0
    code, d = start(repo, "--dry-run")
    assert code == 2 and "ค้างอยู่" in d["error"]                        # dry-run ห้ามทับงานจริงที่ค้าง
    code, d = run_cli(repo, "abort", "--org", "oag", "--month", "2026-09")
    assert code == 0 and d["state"] == "aborted" and d["was"] == "begun"
    assert run_cli(repo, "plan", "--org", "oag", "--month", "2026-09")[0] == 2   # ticket ที่ยกเลิกใช้ต่อไม่ได้
    assert start(repo)[0] == 0                                            # begin ใหม่ได้โดยไม่ต้อง --force


def test_begin_archives_stale_work_files(repo):
    start(repo, "--dry-run")
    feed(repo, [M("m1", "10-09-2026 09:00", "AAA111", 150)])
    code, d = start(repo, "--dry-run")
    assert code == 0 and d["archived_previous_work"]
    assert not (repo / "_work/oag/2026-09/meta.jsonl").exists()
    assert list((repo / "_work/archive").glob("oag-2026-09-*/meta.jsonl"))
    assert feed(repo, [M("m1", "10-09-2026 09:00", "AAA111", 150)])[1]["accepted"] == 1   # ไม่ถูกนับเป็น duplicate ของรอบก่อน


def test_test_env_uses_test_label(repo):
    code, d = start(repo, "--dry-run", "--env", "test")
    assert d["label"] == "zz-test-oag-2026-09"


def test_orgs_are_isolated(repo):
    start(repo, "--dry-run")
    assert run_cli(repo, "status")[1]["tickets"][0]["org"] == "oag"
    code, d = start(repo, "--dry-run", org="mwd")
    assert code == 0 and "01_ประปา" in d["report_file"] and "02_สตง" not in d["report_file"]
    # เมตาของ oag ไม่ปนกับ mwd
    feed(repo, [M("m1", "10-09-2026 09:00", "AAA111", 1)])
    code, d = run_cli(repo, "add-meta", "--org", "mwd", "--month", "2026-09", stdin=jl(M("m1", "10-09-2026 09:00", "XXX111", 1)))
    assert d["rejected_wrong_sender"] == 1 and d["accepted"] == 0       # ผู้ส่งของ oag เข้า mwd ไม่ได้
